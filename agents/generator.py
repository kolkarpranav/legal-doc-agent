import os
import json
from dotenv import load_dotenv
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

import sys
# Ensure we can import from the parent directory
sys.path.append(os.path.join(os.path.dirname(__file__), ".."))
from schemas.entities import CaseEntities
from utils.pdf_reader import extract_text_from_pdf
from utils.llm_client import call_llm

load_dotenv()

FIXED_PHRASES = [
    "do hereby solemnly affirm and state as under:",
    "am well acquainted with the facts and circumstances of the case",
    "I have perused the Petition and the documents annexed thereto",
    "am competent to affirm this Affidavit in Reply",
    "At the outset, I deny each and every allegation, contention and submission",
    "save and except those specifically admitted herein",
    "misconceived, devoid of merits and is liable to be dismissed in limine",
    "strictly in accordance with law and after following due procedure",
    "In the premises aforesaid",
    "I therefore respectfully pray that this Hon'ble Court may be pleased to:",
    "grant such other and further reliefs as this Hon'ble Court may deem fit and proper",
    "true and correct to my knowledge and belief",
    "nothing material has been concealed therefrom"
]

GENERATOR_SYSTEM_PROMPT = """You are an expert legal drafter. Your task is to generate a precise legal document (Affidavit in Reply) based on provided entities and rules.
You MUST strictly follow the provided structure rules and the sample affidavit format.
You MUST include all 13 fixed legal phrases VERBATIM exactly as requested.
Do NOT use markdown (like ```). Output the raw text of the affidavit, maintaining paragraph breaks and clear section headings. Ensure the verification explicitly says "paragraphs 1 to 6"."""

GENERATOR_USER_PROMPT = """
--- FORMAT RULES (From Document 1) ---
{rules_text}

--- SAMPLE AFFIDAVIT (Few-shot example from Document 2) ---
{sample_text}

--- CASE ENTITIES ---
{entities_json}

--- FIXED LEGAL PHRASES REQUIRED (MUST APPEAR VERBATIM) ---
{fixed_phrases}

TASK: Generate the complete Affidavit in Reply. Ensure you include all 10 required sections in order:
1. Forum Heading
2. Jurisdiction
3. Case Number
4. Cause Title
5. Affidavit Title
6. Deponent Clause
7. Numbered Body Paragraphs
8. Prayer (lettered a, b, c)
9. Jurat
10. Verification
Also include the Advocate block.
Ensure paragraph numbers are 1 to 6.
"""

def generate_affidavit(entities_path: str, inputs_dir: str, output_docx_path: str) -> None:
    """Loads entities, constructs few-shot prompt, calls Grok, saves text and formats DOCX."""
    # Load entities
    if not os.path.exists(entities_path):
        raise FileNotFoundError(f"Entities file not found at {entities_path}. Please run Extractor first.")
        
    with open(entities_path, 'r', encoding='utf-8') as f:
        entities_data = json.load(f)
        
    # Parse back into Pydantic model to ensure it's valid
    entities = CaseEntities(**entities_data)
    
    # Read reference PDFs
    rules_path = os.path.join(inputs_dir, "01_Affidavit_Format_Explained.pdf")
    sample_path = os.path.join(inputs_dir, "02_Affidavit_in_Reply_Sample.pdf")
    
    rules_text = extract_text_from_pdf(rules_path) if os.path.exists(rules_path) else "Rules PDF not found."
    sample_text = extract_text_from_pdf(sample_path) if os.path.exists(sample_path) else "Sample PDF not found."
    
    # Construct prompt
    prompt = GENERATOR_USER_PROMPT.format(
        rules_text=rules_text,
        sample_text=sample_text,
        entities_json=entities.model_dump_json(indent=2),
        fixed_phrases="\n".join([f"- {p}" for p in FIXED_PHRASES])
    )
    
    try:
        full_prompt = f"{GENERATOR_SYSTEM_PROMPT}\n\n{prompt}"
        affidavit_text = call_llm(full_prompt).strip()
        
        # Clean markdown if accidentally included
        if affidavit_text.startswith("```"):
            lines = affidavit_text.split('\n')
            if len(lines) > 1 and lines[0].startswith("```"):
                lines = lines[1:]
            if len(lines) > 0 and lines[-1].startswith("```"):
                lines = lines[:-1]
            affidavit_text = "\n".join(lines).strip()

        # STEP 1: Always save the raw text first
        txt_path = output_docx_path.replace(".docx", ".txt")
        os.makedirs(os.path.dirname(txt_path), exist_ok=True)
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(affidavit_text)
        print(f"      Saved: {txt_path}")
            
        # Format docx
        doc = Document()
        for line in affidavit_text.split('\n'):
            line = line.strip()
            if not line:
                continue
                
            p = doc.add_paragraph()
            # Simple heuristic formatting for DOCX styling
            if line.isupper() and len(line) < 100 and "AFFIDAVIT" in line or "COURT" in line or "JURISDICTION" in line or "WRIT PETITION" in line:
                # Center headings in all caps
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line)
                run.bold = True
            elif "DEPONENT" in line and len(line) < 20:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                run = p.add_run(line)
                run.bold = True
            elif line.startswith("Before Me") or line.startswith("BEFORE ME") or line.startswith("Advocates") or "Associates" in line:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(line)
                run.bold = True
            elif line[0].isdigit() and (len(line) > 1 and (line[1] == '.' or line[2] == '.')):
                # Numbered paragraph - bold the number
                parts = line.split('.', 1)
                run = p.add_run(parts[0] + ".")
                run.bold = True
                if len(parts) > 1:
                    p.add_run(parts[1])
            elif line.startswith('(') and len(line) > 2 and line[2] == ')' and line[1].isalpha():
                # Lettered prayer - bold the letter
                parts = line.split(')', 1)
                run = p.add_run(parts[0] + ")")
                run.bold = True
                if len(parts) > 1:
                    p.add_run(parts[1])
            else:
                p.add_run(line)
                
        doc.save(output_docx_path)
        
    except Exception as e:
        print(f"Error in Generator: {e}")
        raise

if __name__ == "__main__":
    entities_path = os.path.join(os.path.dirname(__file__), "..", "outputs", "entities.json")
    inputs_dir = os.path.join(os.path.dirname(__file__), "..", "inputs")
    output_docx = os.path.join(os.path.dirname(__file__), "..", "outputs", "affidavit.docx")
    
    print("Generating affidavit...")
    try:
        generate_affidavit(entities_path, inputs_dir, output_docx)
        print(f"Success! Output saved to {output_docx}")
        print(f"Raw text also saved to {output_docx.replace('.docx', '.txt')}")
    except Exception as e:
        print(f"Generation failed: {e}")
